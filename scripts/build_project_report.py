"""Source-based project report; no credentials or new hardware operations."""
from pathlib import Path
import json
from xml.sax.saxutils import escape
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables/project-report';OUT.mkdir(parents=True,exist_ok=True)
QA=ROOT/'benchmarks/review/project-report';QA.mkdir(parents=True,exist_ok=True)
campaign=json.loads((ROOT/'benchmarks/raw/hardware/ecosystem_realtime_20261022T084316Z.json').read_text())
manifest=json.loads((ROOT/'deliverables/delivery-manifest.json').read_text())

pages=[]
def page(title):
    blocks=[];pages.append((title,blocks));return blocks
def p(b,text):b.append(('p',text))
def h(b,text):b.append(('h',text))
def table(b,head,rows,widths):b.append(('t',(head,rows,widths)))

b=page('ClawGrid Home Automation Project Report')
p(b,'ESP32 S3 control with offline PC voice and a shared Supabase ecosystem')
p(b,'Technical project report | 22 October 2026 | Review of supplied archives, current source and retained verification evidence')
h(b,'Abstract')
p(b,'ClawGrid is a two-board home automation prototype for a bedroom light and fan. Remote ESP32-S3 Node A reads one passive infrared sensor and owns deterministic device policy. PC-connected Node B bridges commands and telemetry over authenticated ESP-NOW. A local PC recognizes speech offline and connects the boards to a webpage, an Android client and Supabase status and command history. The design keeps the final decision to accept an action on A, including when a command originates from speech or a cloud client. [S1-S6]')
p(b,'The software ecosystem has been implemented and packaged. Retained hardware evidence shows encrypted Supabase commands reaching A through B, with matching application acknowledgements and virtual ON/OFF state. The last recorded hardware condition is a Windows USB error on B COM7, so those earlier results do not establish a currently healthy connection. Physical relay GPIOs remain inhibited, and real load switching has not been validated. [S7-S10]')
h(b,'Scope and evidence boundary')
p(b,'This report assesses the current ESP/PC project and the two supplied archives: ClawGrid-frontend-updated rtos.zip and clawgrid.zip. The original archives remain intact. Working copies are under ecosystem/web-baseline and ecosystem/android. The report reviews saved files and logs; it does not claim a new live bench test during report preparation. No student name, institution, completion percentage or experimental accuracy has been invented.')
table(b,['Area','Recorded outcome'],[
 ['Gateway and cloud','Implemented; retained encrypted command test passed'],
 ['Web and Android clients','Shared protocol; APK build and host wire tests passed'],
 ['Latest USB condition','B COM7 unavailable with Windows device error 31'],
 ['Physical loads','Not tested; normal relay outputs inhibited'],
 ['Sensor coverage','One bedroom PIR; temperature and lux unavailable']], [1.65,4.95])

b=page('Project objectives and architecture')
h(b,'Problem and objectives')
p(b,'The project aims to accept natural spoken light/fan requests, transfer explicit commands to a remote controller, and show the same board observations in a local webpage and an Android app. A user should be able to say bedroom light on, later say bedroom light off, and inspect a controller acknowledgement without assuming that a lamp contact or electrical load has been independently measured. Continuous speech capture is user-started and remains under the PC page controls. [S1-S4]')
p(b,'The engineering objectives are to preserve deterministic control on the ESP, distinguish physical and virtual rooms, expose fresh measured telemetry, prevent stale or duplicate commands from executing, and keep private cloud credentials outside client packages. The PC remains a required gateway for cloud operation; this version is not an autonomous ESP-to-cloud deployment.')
b.append(('image',str(QA/'architecture.png')))
table(b,['Component','Responsibility'],[
 ['Node A','PIR sampling; occupancy; light/fan policy; SafetyKernel; command ACKs'],
 ['Node B','USB API tunnel; proposals; authenticated radio forwarding; own diagnostics'],
 ['PC gateway','Offline Vosk speech; serial transport; Supabase sync and encrypted snapshots'],
 ['Supabase','Broadcast transport, durable telemetry and command receipts'],
 ['Web and Android','Decode shared observations; explicit commands; stale/unknown status']], [1.35,5.25])
p(b,'A is the relay/PIR board, assigned COM8 when programming. B is the bare PC bridge on COM7. Their recorded MAC assignments end f5:bc:b0 and f6:a9:68 respectively. The native build names are historical: A uses native_voice, while B uses native_intelligence. Speech recognition in the deployed workflow runs on the PC. A and B communicate wirelessly; no inter-board TX/RX cable is required. [S1, S2]')

b=page('Firmware control and offline voice')
h(b,'Task ownership')
p(b,'The native entry point selects each board role and starts the role implementation. On A, ControlTask owns the safety registry after startup, samples PIR input, advances automation and handles bounded command work. LinkTask transfers authenticated messages and reports. B runs its link and command bridge task. A low-priority Diagnostics task runs separately; a traffic LED task handles visual activity without placing LED timing in radio callbacks. [S3-S5]')
table(b,['Task or mechanism','Source configuration'],[
 ['A ControlTask','Core 1; priority 5; 5 ms control period'],
 ['A LinkTask','Core 0; priority 3; static task and stack allocation'],
 ['B LinkTask','Core 0; priority 3; measured API reports 5 ms period'],
 ['Command and ACK mailboxes','A uses bounded static queues with eight entries each'],
 ['State publication','Latest-value mailbox avoids an accumulating state backlog']], [2.2,4.4])
h(b,'Safety and automation')
p(b,'SafetyKernel defines a 3000 ms minimum dwell between transitions, a default five-minute manual hold and a four-hour maximum continuous run. These are source defaults, not a claim that every limit was physically exercised in the latest electrical setup. Automation supports device requests, scenes, timers, vacancy handling, manual overrides and saved policy. The default source policy enables automatic lighting and leaves automatic fan operation off; saved policy can override defaults. [S3, S5]')
p(b,'The PIR logic allows a 60-second warm-up, requires a LOW baseline, and filters stable changes over 200 ms according to the project wiring guide. One sensor supplies evidence for its bedroom zone only. Temperature-based fan control and independent room presence cannot be claimed from this hardware. Normal firmware keeps physical GPIO switching inhibited pending interface validation. [S1, S2]')
h(b,'Speech path')
p(b,'The /voice page starts microphone capture only when the user chooses Start listening. The local Vosk model processes speech on the PC; the board bridge receives command text rather than microphone audio. Complete recognized commands use the existing proposal and confirmation path. Stop, target changes and bounded upload queues prevent an old audio backlog from becoming a later command stream. Audio is not saved in the implemented path. Human speech accuracy, noise tolerance and sustained microphone operation require a separate recorded evaluation. [S4, S10]')

b=page('Communication and cloud integration')
h(b,'Command sequence')
p(b,'A command client creates an explicit target, room, device and ON/OFF action, with a UUID and a ten-second cloud expiry. The PC stores a signed request in Supabase, checks scope and freshness, obtains B\'s immutable proposal, confirms the exact target, and observes A\'s matching application ACK. A still evaluates the request through its SafetyKernel. A radio delivery indication alone is insufficient to prove that the application accepted the request. [S6]')
p(b,'A local SQLite journal claims a command before execution. Duplicate requests, gateway restarts and reconnects do not automatically repeat an actuation. A missing ACK leaves the outcome unknown. Stored receipts can be uploaded again idempotently, without re-sending the command. Realtime requests are processed after durable storage instead of waiting for the telemetry history cycle. [S6, S10]')
h(b,'Security boundaries')
p(b,'The ESP-NOW link uses encrypted unicast peers and a mandatory application authentication tag. The source appends a 16-byte tag derived from HMAC-SHA256 and checks the expected peer. Cloud Broadcast carries AES-256-GCM envelopes whose key derives from the existing ClawGrid control token. Home and gateway identifiers are authenticated as additional data. Python, browser WebCrypto and Android Kotlin implementations share that wire format. [S5, S6]')
p(b,'The public clients contain the Supabase URL, anon key and public scope identifiers. The service-role key stays in an ignored PC configuration file. Existing database tables and row-level security were preserved; the integration did not grant anonymous table read/write access. The control token is held in client memory. The service-role credential previously exposed in the conversation should be rotated and its local configuration updated. No credential is reproduced in this report. [S1, S6, S9]')
table(b,['Stored object','Purpose'],[
 ['clawgrid_homes and clawgrid_rooms','Home and zone identity; physical capability metadata'],
 ['clawgrid_devices','Device mapping and low-voltage capability metadata'],
 ['clawgrid_telemetry','Bedroom observations with source, known flag and age'],
 ['clawgrid_events','Signed command requests and software ACK receipts'],
 ['clawgrid_policy_snapshots','Existing policy history schema retained']], [2.35,4.25])
p(b,'The archive review found seeded room readings and local simulation in the original app path, plus a service key in its environment. The new live launcher and website entry use shared gateway observations. Legacy demo sources remain available for reference and are not authoritative physical sensor views. [S1, S8]')

b=page('Real data mapping and freshness')
p(b,'The live views separate controller state, measured diagnostics and unavailable signals. HARDWARE_MEASURED labels board observations, HOST_MEASURED labels PC audit activity, and SIMULATED identifies virtual input. A simulated-room command can be processed by real hardware while its room inputs remain synthetic. This distinction is preserved in the retained test evidence. [S1, S7]')
table(b,['Field or view','Origin and interpretation'],[
 ['Bedroom occupancy and motion','A PIR on GPIO6 and its occupancy policy; one physical zone'],
 ['Light and fan state','A\'s commanded state; no independent load/contact feedback'],
 ['A boot uptime and radio counters','Authenticated reports returned through B /api/state'],
 ['B internal heap and PSRAM','B /api/memory; bytes and diagnostic timestamp'],
 ['B task timing and watchdog','B /api/rtos; LinkTask metrics only'],
 ['A heap and task metrics','Not exposed by the current radio protocol; UNKNOWN'],
 ['Temperature and illuminance','No installed sensor; null and not installed'],
 ['Other rooms','No physical sensors assigned; explicitly virtual only'],
 ['Live history','Last 60 received samples in client memory'],
 ['Durable history','Supabase telemetry rows and command receipts']], [2.05,4.55])
h(b,'Freshness contract')
table(b,['Interval or limit','Meaning'],[
 ['About 1 second','Gateway snapshot publishing target'],
 ['About 5 seconds','B diagnostic refresh target'],
 ['About 10 seconds','Durable telemetry upload target'],
 ['Under 3000 ms','A must be known and authenticated within this report age'],
 ['4 seconds without a snapshot','Client removes live state and marks data stale'],
 ['10 seconds for diagnostics','Older B diagnostic values become unavailable']], [2.25,4.35])
p(b,'These intervals are configured targets, not guaranteed end-to-end deadlines. USB, radio and internet delays can increase them. Clients reject old, future, duplicate, out-of-order and wrong-scope snapshots. Command buttons require fresh authenticated A evidence. [S6, S8]')
p(b,'The existing SQL schema requires a non-null outputs_enabled boolean. For known=false rows, its conservative false placeholder must not be displayed as a new hardware observation. Consumers must gate the entire row on known and freshness. Light, fan, motion and absent sensor readings remain null when unavailable. Expected MAC labels come from role configuration and must not be described as a second independently measured MAC field. [S1, S11]')

b=page('Hardware mapping and bench wiring')
p(b,'The mapping below records the intended low-voltage bench connections in the project guide. It does not certify the present physical wiring, module input compatibility or electrical OFF behaviour. The normal image continues to inhibit relay GPIO outputs. [S2]')
table(b,['Connection','Recorded mapping or requirement'],[
 ['A USB power','Powers the remote controller; COM8 used when programming'],
 ['B USB data and power','PC bridge on COM7; both boards must remain powered'],
 ['PIR OUT','A GPIO6 after confirming the actual output is 3.3 V logic'],
 ['PIR VCC and GND','Labeled 5 V rail and GND per the project guide; verify the unit labels'],
 ['Light relay input','A GPIO4 through a verified compatible interface'],
 ['Fan relay input','A GPIO5 through a verified compatible interface'],
 ['Relay coil power','Regulated 5 V relay supply; never ESP 3V3'],
 ['Ground reference','Common reference only when required by the verified non-isolated interface'],
 ['Relay contacts','For an approved DC bench load: fused supply to COM, NO to load, NC unused']], [1.9,4.7])
h(b,'Interface constraints')
p(b,'SRD-05VDC-SL-C identifies a 5 V relay coil. That marking does not establish whether the module IN terminal can safely connect to an ESP GPIO. The full PCB, control header, transistor/opto circuitry, jumper arrangement and input thresholds need inspection. A module with isolated JD-VCC wiring needs its own verified connection plan. The provisional active-low software assumption must match the complete interface, including any inversion. [S2]')
p(b,'The guide advises against joining a separate relay supply positive rail to a USB-powered ESP rail. Reference grounds are joined only where the chosen interface requires them. A defined electrical OFF state must be checked during startup, reset, loss of ESP power and USB reconnect. Logical OFF and a successful application ACK do not demonstrate that the physical contacts are open. Mains AC switching is outside this prototype. [S2, S12]')
h(b,'Deployment sequence')
p(b,'First restore a reliable data connection to bare B. Confirm its identity and assigned role before any future upload. Keep A powered, start the PC bridge on the detected B serial port, and connect the live webpage or app with the existing control token. Validate virtual ON/OFF and acknowledgements first. Only after the relay interface, voltage levels, polarity and reset behaviour are independently verified should an explicit A bench image be considered for unloaded contact tests, followed by suitable low-voltage DC loads. The current report does not authorize or claim those tests. [S1, S2]')

b=page('Verification results and measurements')
p(b,'The following outcomes are taken from retained evidence. They describe the captured campaign and host checks, not continuous operation or the later USB fault. Hardware command testing used the virtual target and left physical GPIOs inhibited. [S7-S11]')
table(b,['Check','Result and interpretation'],[
 ['Cloud command round trip','Virtual light ON and OFF reached A; phase 2 status 0 ACKs; state matched'],
 ['Independent cloud subscriber','Three encrypted snapshots with A state and B diagnostics'],
 ['Gateway and USB regression suite','24 host tests passed'],
 ['Android wire implementation','Three host JUnit tests passed using actual Kotlin/JCE code'],
 ['Browser and Python encryption','Cross-decryption passed; wrong key, scope and tamper rejected'],
 ['Exact web client against WSS','Received encrypted snapshots; correctly represented unavailable hardware'],
 ['Android build and package','Debug APK built; package and signature verification passed'],
 ['Actual relay and phone operation','Not tested; no phone installation performed']], [2.25,4.35])
h(b,'Captured hardware values')
nb=campaign['snapshots'][-1]['node_b'];mem=nb['memory'];rt=nb['rtos']
table(b,['Observation','Captured value'],[
 ['B internal heap free',f"{mem['internal_free']:,} bytes"],
 ['B PSRAM free and total',f"{mem['psram_free']:,} / {mem['psram_total']:,} bytes"],
 ['B LinkTask maximum execution',f"{rt['max_exec_us']:,} us within retained counter history"],
 ['B LinkTask maximum lateness',f"{rt['max_late_us']:,} us within retained counter history"],
 ['B deadline misses',str(rt['deadline_misses'])+' in the captured counter'],
 ['ON radio transaction RTT','85.196 ms; one radio attempt'],
 ['OFF radio transaction RTT','45.257 ms; one radio attempt'],
 ['A command processing','25 us ON and 28 us OFF for these virtual requests']], [3.5,3.1])
p(b,'The ON/OFF timing sample size is two commands. The radio RTT fields are not speech latency, internet latency or complete browser-to-relay response time. No percentile, average service-level guarantee, recognition accuracy, power saving, radio range or relay reliability result is established by this sample. The diagnostic maxima are counters in one saved snapshot; their observation duration and electrical switching conditions were not controlled as a benchmark. [S7]')

b=page('Limitations and completion plan')
h(b,'Latest recorded blockers')
p(b,'After the successful command campaign, B\'s native USB stopped returning application bytes. Normal reconnect/reset and a manual BOOT-mode identity check did not restore it. The retained log reports Windows error 31, device attached not functioning, while configuring COM7. A Windows device restart attempt was denied by OS permissions. The PC service was adjusted to keep the website and encrypted cloud status available with unknown hardware data, rather than terminate or fabricate readings. No firmware was erased during recovery. This report does not establish that the fault is in the firmware, cable or Windows driver. [S1, S11]')
p(b,'The other completion limits are physical: relay input compatibility and power wiring remain unverified; physical outputs are inhibited; only one physical PIR is available; temperature and lux are absent; A\'s heap/task metrics are not exposed remotely; and there is no independent contact/load feedback. The Android build and crypto tests do not substitute for an on-device network and lifecycle test. Speech accuracy and long-duration operation also remain unmeasured.')
h(b,'Next validation steps')
table(b,['Priority','Action and acceptance evidence'],[
 ['1','Test B with a known-working USB data cable and direct PC port; release BOOT and RESET. Require stable application replies and fresh authenticated A reports.'],
 ['2','Repeat virtual ON/OFF through web and Android; retain matching UUIDs, A ACKs and state; verify no command replay after reconnect.'],
 ['3','Verify relay module input circuitry, logic levels, power rails, polarity and reset OFF state. Keep loads disconnected until verified.'],
 ['4','Perform approved low-voltage contact/load tests and PIR warm-up, detection and vacancy tests; record independent observations.'],
 ['5','Measure speech command accuracy and full-path latency with a stated corpus, trial count, noise conditions and distribution.'],
 ['6','Run sustained disconnect, restart, timer, dwell, watchdog and memory tests; rotate the exposed service credential.']], [0.65,5.95])
h(b,'Project assessment')
p(b,'ClawGrid has an implemented software path joining the PC, both ESP roles, Supabase and shared clients, with evidence that real hardware processed cloud commands for an explicit virtual target. It remains a prototype requiring restoration of B\'s USB link and controlled electrical validation before physical relay operation can be claimed. The source and retained results support deterministic control, bounded command handling and honest data provenance; they do not establish production readiness or a research novelty claim.')

b=page('Deliverables and source references')
p(b,'All paths below are relative to ClawGrid-physical-integration-prep. Private credentials, radio keys and control-token values are deliberately omitted. The originals in Downloads remain unchanged. The local report is an assessment of this project snapshot, not a replacement for the source and logs.')
table(b,['Delivered component','Location'],[
 ['Android APK','deliverables/ClawGrid-Live-debug.apk'],
 ['Android source','deliverables/ClawGrid-Android-source.zip'],
 ['Standalone website','deliverables/ClawGrid-live-web.zip'],
 ['Gateway and cloud source','scripts/browser_bridge.py, supabase_bridge.py, ecosystem_realtime.py'],
 ['Main firmware sources','firmware/node_voice/src/main.cpp and firmware/node_intelligence/src/main.cpp'],
 ['Package hashes and scope','deliverables/delivery-manifest.json']], [1.7,4.9])
refs=[
 ('S1','Integration scope and exact mapping','docs/ECOSYSTEM_INTEGRATION.md'),
 ('S2','Roles and provisional bench wiring','docs/PROGRAM_AND_WIRE.md'),
 ('S3','A control ownership and B bridge','firmware/node_voice/src/main.cpp; firmware/node_intelligence/src/main.cpp'),
 ('S4','PC speech and continuous capture','scripts/offline_speech.py; firmware/node_intelligence/src/dashboard_client.h'),
 ('S5','Native tasks and deterministic control','native/main/app_entry.cpp; firmware/common/safety_kernel/safety_kernel.h; firmware/common/runtime/automation.h; secure_link.h'),
 ('S6','Cloud command and wire implementation','scripts/supabase_bridge.py; scripts/ecosystem_realtime.py; dashboard/ecosystem-client.js'),
 ('S7','Successful retained hardware campaign','benchmarks/raw/hardware/ecosystem_realtime_20261022T084316Z.json'),
 ('S8','Android shared live client','ecosystem/android/app/src/main/java/com/example/EcosystemActivity.kt; data/remote/EcosystemClient.kt; RealtimeWire.kt'),
 ('S9','Delivered package verification','deliverables/delivery-manifest.json; benchmarks/raw/android-delivery-build.log'),
 ('S10','Host tests and Android wire tests','benchmarks/raw/supabase_unit_20261022_final.xml; benchmarks/raw/android-wire-jvm.log; scripts/check_ecosystem_crypto.cjs'),
 ('S11','Cloud history and disconnected state','benchmarks/raw/hardware/ecosystem_history_and_disconnect_20261022.json; benchmarks/raw/ecosystem_web_client_live.json; bridge_manual_boot_identity_20261022.log in benchmarks/raw/hardware'),
 ('S12','Prototype constraints','AGENTS.md')]
for ref,title,path in refs:p(b,f'[{ref}] {title}. {path}')

# A simple shared architecture figure with readable vector-like boxes.
im=Image.new('RGB',(1500,540),'white');dr=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',30)
small=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',23)
def box(x,y,w,title,sub):
    dr.rounded_rectangle((x,y,x+w,y+110),radius=12,fill='#F2F5F7',outline='#58636B',width=2)
    dr.text((x+w/2,y+33),title,font=font,fill='black',anchor='mm');dr.text((x+w/2,y+76),sub,font=small,fill='#33414C',anchor='mm')
def arrow(x1,y1,x2,y2):
    dr.line((x1,y1,x2,y2),fill='#33414C',width=4)
    if y1==y2:dr.polygon([(x2,y2),(x2-14,y2-8),(x2-14,y2+8)],fill='#33414C')
    else:dr.polygon([(x2,y2),(x2-8,y2-14),(x2+8,y2-14)],fill='#33414C')
box(30,35,360,'Remote Node A','PIR and SafetyKernel');box(555,35,330,'USB Node B','Command bridge');box(1060,35,405,'PC gateway','Offline voice and cloud sync')
arrow(390,90,555,90);dr.text((470,37),'ESP-NOW',font=small,fill='black',anchor='mm')
arrow(885,90,1060,90);dr.text((972,37),'USB',font=small,fill='black',anchor='mm')
box(1060,280,405,'Supabase','Encrypted Broadcast and history');arrow(1262,145,1262,280)
box(30,280,360,'Webpage','Shared live observations');box(555,280,330,'Android','Shared live observations')
dr.line((1060,335,885,335),fill='#33414C',width=4);dr.polygon([(885,335),(899,327),(899,343)],fill='#33414C')
dr.line([(965,335),(965,220),(210,220),(210,280)],fill='#33414C',width=4)
dr.polygon([(210,280),(202,266),(218,266)],fill='#33414C')
dr.text((750,470),'Commands return through the same gateway and A safety checks',font=small,fill='black',anchor='mm')
im.save(QA/'architecture.png')

doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
for name,size in [('Normal',10),('Title',25),('Heading 1',19),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.08
doc.styles['Title'].paragraph_format.space_after=Pt(12)
header=sec.header.paragraphs[0];header.text='CLAWGRID   |   PROJECT TECHNICAL REPORT';header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('22 October 2026  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for i,(title,blocks) in enumerate(pages):
    if i:doc.add_page_break()
    doc.add_paragraph(title,'Title' if i==0 else 'Heading 1')
    for kind,value in blocks:
        if kind=='p':doc.add_paragraph(value)
        elif kind=='h':doc.add_paragraph(value,'Heading 2')
        elif kind=='image':doc.add_picture(value,width=Inches(6.7))
        else:
            head,rows,widths=value;t=doc.add_table(rows=1,cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
            for j,text in enumerate(head):t.rows[0].cells[j].text=text
            repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows:
                for cell,text in zip(t.add_row().cells,row):cell.text=text
            for ri,row in enumerate(t.rows):
                for j,cell in enumerate(row.cells):
                    cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'303A43' if ri==0 else ('F3F5F7' if ri%2==0 else 'FFFFFF'));pr.append(shade)
                    margins=OxmlElement('w:tcMar')
                    for edge in ('top','left','bottom','right'):
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
                    pr.append(margins)
                    borders=OxmlElement('w:tcBorders')
                    for edge in ('top','left','bottom','right'):
                        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    pr.append(borders)
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.04
                        for run in para.runs:run.font.size=Pt(9);run.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
doc.save(OUT/'ClawGrid_Project_Report.docx')

styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='BodyReport',fontName='Helvetica',fontSize=10,leading=13.5,spaceAfter=8))
styles.add(ParagraphStyle(name='CellReport',fontName='Helvetica',fontSize=8.7,leading=11))
styles.add(ParagraphStyle(name='HeaderCellReport',fontName='Helvetica-Bold',fontSize=9,leading=11,textColor=colors.white))
styles['Title'].fontName='Helvetica-Bold';styles['Title'].fontSize=25;styles['Title'].leading=30;styles['Title'].textColor=colors.black;styles['Title'].alignment=0
styles['Heading1'].textColor=colors.black;styles['Heading1'].fontSize=18;styles['Heading1'].leading=22
styles['Heading2'].textColor=colors.black;styles['Heading2'].fontSize=12;styles['Heading2'].leading=15
story=[]
for i,(title,blocks) in enumerate(pages):
    if i:story.append(PageBreak())
    story.append(Paragraph(escape(title),styles['Title' if i==0 else 'Heading1']))
    story.append(Spacer(1,8))
    for kind,value in blocks:
        if kind in ('p','h'):story.append(Paragraph(escape(value),styles['BodyReport' if kind=='p' else 'Heading2']))
        elif kind=='image':story.extend([RLImage(value,width=483,height=174),Spacer(1,8)])
        else:
            head,rows,widths=value;data=[[Paragraph(escape(x),styles['HeaderCellReport']) for x in head]]+[[Paragraph(escape(str(x)),styles['CellReport']) for x in row] for row in rows]
            t=Table(data,colWidths=[x/sum(widths)*483 for x in widths],repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#303A43')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F3F5F7')]),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
            story.extend([t,Spacer(1,10)])
def footer_pdf(canvas,document):
    canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#44505B'))
    canvas.drawString(50,810,'CLAWGRID  |  PROJECT TECHNICAL REPORT')
    canvas.drawRightString(545,27,f'22 October 2026  |  {document.page}')
SimpleDocTemplate(str(OUT/'ClawGrid_Project_Report.pdf'),pagesize=A4,rightMargin=55,leftMargin=55,topMargin=55,bottomMargin=48,title='ClawGrid Home Automation Project Report',author='ClawGrid project review').build(story,onFirstPage=footer_pdf,onLaterPages=footer_pdf)
md=[]
for title,blocks in pages:
    md+=['# '+title,'']
    for kind,value in blocks:
        if kind=='p':md+=[value,'']
        elif kind=='h':md+=['## '+value,'']
        elif kind=='t':
            head,rows,_=value;md+=['| '+' | '.join(head)+' |','| '+' | '.join(['---']*len(head))+' |',*['| '+' | '.join(row)+' |' for row in rows],'']
(OUT/'ClawGrid_Project_Report.md').write_text('\n'.join(md),encoding='utf-8')
print('Created report PDF DOCX and Markdown:',OUT)
